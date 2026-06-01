#!/usr/bin/env python3
"""Strict extractor: reads the Career Framework spreadsheet and emits a
structured JSON the importer loads verbatim. Mapping is per-sheet because each
sheet has a slightly different column layout.

Run: python3 scripts/extract_framework.py <xlsx> prisma/framework-content.json
"""
import json
import sys
import openpyxl

# ── Level codes / labels (from the sheet headers) ───────────────────────────
IC_CODES = ["IC2", "IC3", "IC4", "IC5"]
M_CODES = ["M4", "M5", "M6"]

LEVEL_LABELS = {
    "IC2": "Beginner",
    "IC3": "Proficient",
    "IC4": "Fully proficient",
    "IC5": "Domain expert",
    "M4": "Team Lead",
    "M5": "Senior Manager",
    "M6": "Strategic Leader",
}


def clean_name(v):
    if v is None:
        return ""
    # collapse internal newlines/whitespace in names
    return " ".join(str(v).split()).strip()


def clean_text(v):
    if v is None:
        return None
    s = " ".join(str(v).split()).strip()
    return s or None


def split_bullets(v):
    """Split a cell into an ordered array of bullet strings, stripping the
    leading bullet marker (•, -, *, –) and surrounding whitespace."""
    if v is None:
        return []
    out = []
    for line in str(v).split("\n"):
        s = line.strip()
        if not s:
            continue
        # strip a leading bullet marker
        while s[:1] in ("•", "-", "*", "–", "·"):
            s = s[1:].strip()
        s = s.strip()
        if s:
            out.append(s)
    return out


def extract_competencies(ws, name_col, desc_col, level_map, start_row, end_row):
    """Generic row extractor. A row is a competency when its name cell is
    non-empty AND at least one level cell has content."""
    rows = []
    for r in range(start_row, end_row + 1):
        name = clean_name(ws.cell(r, name_col).value)
        if not name:
            continue
        levels = {}
        for code, col in level_map.items():
            bullets = split_bullets(ws.cell(r, col).value)
            if bullets:
                levels[code] = bullets
        if not levels:
            continue
        comp = {"name": name, "levels": levels}
        if desc_col:
            d = clean_text(ws.cell(r, desc_col).value)
            if d:
                comp["description"] = d
        rows.append(comp)
    return rows


def main():
    src = sys.argv[1]
    out = sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=True)

    # ── General competencies: merge IC + M sheets by competency name ──────────
    ic_ws = wb["INDIVIDUAL CONTRIBUTORS"]
    m_ws = wb["MANAGERS"]

    ic_general = extract_competencies(
        ic_ws, 1, None, {"IC2": 2, "IC3": 3, "IC4": 4, "IC5": 5}, 2, 6
    )
    m_general = extract_competencies(
        m_ws, 1, None, {"M4": 2, "M5": 3, "M6": 4}, 2, 6
    )
    m_by_name = {c["name"]: c for c in m_general}

    general = []
    for c in ic_general:
        merged = {"name": c["name"], "levels": dict(c["levels"])}
        m = m_by_name.get(c["name"])
        if m:
            merged["levels"].update(m["levels"])
        general.append(merged)

    # ── Functional sheets: (sheet, family display name, mapping) ──────────────
    families_cfg = [
        ("Engineering", "Engineering", 1, None, {"IC2": 2, "IC3": 3, "IC4": 4, "IC5": 5}, 2, 60),
        ("Product Management", "Product Management", 1, 2, {"IC2": 3, "IC3": 4, "IC4": 5, "IC5": 6}, 2, 60),
        ("Product Design", "Product Design", 2, 3, {"IC2": 4, "IC3": 5, "IC4": 6, "IC5": 7}, 2, 60),
        ("Customer Support", "Customer Support", 1, None, {"IC2": 2, "IC3": 3, "IC4": 4, "IC5": 5}, 2, 60),
        ("Marketing", "Marketing", 1, 2, {"IC2": 3, "IC3": 4, "IC4": 5, "IC5": 6}, 2, 60),
        ("Sales", "Sales", 5, 6, {"IC2": 7, "IC3": 8, "IC4": 9, "IC5": 10}, 2, 60),
        ("Talent Acq.", "Talent Acquisition", 1, 2, {"IC2": 3, "IC3": 4, "IC4": 5, "IC5": 6}, 2, 60),
        ("WIP - Accounting", "Accounting", 1, 2, {"IC2": 3, "IC3": 4, "IC4": 5, "IC5": 6}, 3, 60),
    ]

    families = []
    for order, (sheet, family_name, name_col, desc_col, level_map, start, end) in enumerate(
        families_cfg, start=1
    ):
        ws = wb[sheet]
        comps = extract_competencies(ws, name_col, desc_col, level_map, start, end)
        families.append(
            {
                "name": family_name,
                "sourceSheet": sheet,
                "displayOrder": order,
                "competencies": comps,
            }
        )

    data = {"levelLabels": LEVEL_LABELS, "general": general, "families": families}
    with open(out, "w") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    # summary
    print(f"General competencies: {len(general)}")
    for g in general:
        print(f"  - {g['name']}: levels {sorted(g['levels'].keys())}")
    print("Families:")
    for fam in families:
        print(f"  - {fam['name']} ({fam['sourceSheet']}): {len(fam['competencies'])} competencies")
        for c in fam["competencies"]:
            nb = sum(len(v) for v in c["levels"].values())
            print(f"      · {c['name']}  [{len(c['levels'])} levels, {nb} bullets]")


if __name__ == "__main__":
    main()
